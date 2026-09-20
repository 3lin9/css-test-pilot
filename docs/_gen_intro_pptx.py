"""Generate TestPilot intro PPT for leadership / colleagues."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

# Widescreen 16:9
W, H = Inches(13.333), Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY2 = RGBColor(0x12, 0x2B, 0x4A)
TEAL = RGBColor(0x1A, 0xA3, 0xA8)
TEAL_DK = RGBColor(0x0E, 0x7A, 0x7E)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFF = RGBColor(0xF5, 0xF7, 0xFA)
INK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5B, 0x67, 0x7A)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD7, 0xDE, 0xE8)
SOFT = RGBColor(0xE8, 0xF4, 0xF5)


def set_run(run, size=18, bold=False, color=INK, font="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)


def add_text(tf, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, space_after=6):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    if not tf.paragraphs[0].text and tf.paragraphs[0] is p:
        pass
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def round_box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.adjustments[0] = 0.08
    sh.shadow.inherit = False
    return sh


def tb(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(l, t, w, h)
    sh.text_frame.word_wrap = True
    sh.text_frame.auto_size = None
    try:
        sh.text_frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = sh.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return sh


def multi(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color, space_after)"""
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr"}[anchor])
    except Exception:
        pass
    first = True
    for item in lines:
        text, size, bold, color = item[:4]
        sa = item[4] if len(item) > 4 else 8
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(sa)
        run = p.add_run()
        run.text = text
        set_run(run, size, bold, color)
    return sh


PAGE_MARKS = []


def footer(slide, *_ignored):
    """页码在全部页面生成后统一回填,插页时无需手工改编号。"""
    box(slide, 0, Inches(7.22), W, Inches(0.28), NAVY)
    tb(slide, Inches(0.4), Inches(7.22), Inches(8), Inches(0.28), "TestPilot  ·  AI Native 跨端业务测试基础设施  ·  内部介绍", 10, False, RGBColor(0xA8, 0xC5, 0xC8), PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    mark = tb(slide, Inches(11.2), Inches(7.22), Inches(1.7), Inches(0.28), "", 10, False, WHITE, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    PAGE_MARKS.append((len(prs.slides._sldIdLst), mark))


def node(slide, l, t, w, h, text, fill=WHITE, color=NAVY, size=13, line=LINE):
    """流程图节点"""
    round_box(slide, l, t, w, h, fill, line)
    tb(slide, l + Inches(0.06), t, w - Inches(0.12), h, text, size, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def diamond(slide, l, t, w, h, text, fill=GOLD, color=WHITE, size=12):
    """判断节点"""
    sh = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    tb(slide, l, t, w, h, text, size, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def _arrow(slide, shape, l, t, w, h, fill=TEAL):
    sh = slide.shapes.add_shape(shape, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def ar_r(slide, l, t, w=Inches(0.3), h=Inches(0.18), fill=TEAL):
    return _arrow(slide, MSO_SHAPE.RIGHT_ARROW, l, t, w, h, fill)


def ar_l(slide, l, t, w=Inches(0.3), h=Inches(0.18), fill=TEAL):
    return _arrow(slide, MSO_SHAPE.LEFT_ARROW, l, t, w, h, fill)


def ar_d(slide, l, t, w=Inches(0.18), h=Inches(0.3), fill=TEAL):
    return _arrow(slide, MSO_SHAPE.DOWN_ARROW, l, t, w, h, fill)


def ar_label(slide, l, t, text, color=NAVY):
    tb(slide, l, t, Inches(0.7), Inches(0.24), text, 11, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def header_bar(slide, title, subtitle=None):
    box(slide, 0, 0, W, Inches(1.12), NAVY)
    box(slide, 0, Inches(1.12), W, Inches(0.06), TEAL)
    tb(slide, Inches(0.5), Inches(0.18), Inches(11), Inches(0.5), title, 26, True, WHITE)
    if subtitle:
        tb(slide, Inches(0.5), Inches(0.68), Inches(12), Inches(0.32), subtitle, 12, False, RGBColor(0x9E, 0xD4, 0xD6))


def content_bg(slide):
    box(slide, 0, 0, W, H, OFF)


def card(slide, l, t, w, h, title, body, accent=TEAL):
    sh = round_box(slide, l, t, w, h, WHITE, LINE)
    box(slide, l, t, Inches(0.08), h, accent)
    tb(slide, l + Inches(0.22), t + Inches(0.12), w - Inches(0.35), Inches(0.38), title, 15, True, NAVY)
    tb(slide, l + Inches(0.22), t + Inches(0.5), w - Inches(0.35), h - Inches(0.62), body, 12, False, MUTED)
    return sh


TOTAL = 29
prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# ========== 1 Cover ==========
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, NAVY)
box(s, 0, 0, Inches(0.18), H, TEAL)
tb(s, Inches(0.7), Inches(1.35), Inches(11), Inches(0.4), "内部技术分享  ·  V0.1", 14, False, TEAL)
tb(s, Inches(0.7), Inches(1.85), Inches(12), Inches(1.1), "TestPilot", 48, True, WHITE)
tb(s, Inches(0.7), Inches(2.95), Inches(12), Inches(0.7), "AI Native 跨端业务测试基础设施", 26, False, RGBColor(0xC5, 0xE8, 0xEA))
tb(
    s,
    Inches(0.7),
    Inches(3.8),
    Inches(11.5),
    Inches(1.0),
    "让 AI 能理解、生成、校验并执行业务测试。\n不是又一个 Playwright 封装，而是一套可被人和 Agent 共同使用的测试语言与执行体系。",
    16,
    False,
    RGBColor(0xB8, 0xC4, 0xD4),
)
# bottom chips
for i, t in enumerate(["DSL 共用", "Adapter 可插拔", "Evidence-first", "Git 为真相源", "SDK 统一入口"]):
    x = Inches(0.7) + Inches(i * 2.4)
    round_box(s, x, Inches(5.35), Inches(2.2), Inches(0.48), NAVY2, TEAL)
    tb(s, x, Inches(5.35), Inches(2.2), Inches(0.48), t, 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
tb(s, Inches(0.7), Inches(6.5), Inches(11.5), Inches(0.35), "面向领导与同事  ·  重点讲清：项目结构各层干什么、DSL 是啥、Adapter 是啥", 13, False, RGBColor(0x7A, 0x8E, 0xA8))


# ========== 2 Agenda ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "今天要讲清的六件事", "尤其要听懂三样：仓库每层干什么、DSL 是什么、Adapter 是什么（配 5 张流程图）")
footer(s)
items = [
    ("01", "定位与边界", "工具 vs 业务测试项目；我们管什么、不管什么"),
    ("02", "DSL 是什么", "人和 AI 共用的测试说明书（YAML），不是业务代码"),
    ("03", "Adapter 是什么", "把说明书翻译成「真的去点页面」的驱动插件"),
    ("04", "五张流程图", "端到端 / 单次 Run 判断 / AI 闭环 / init / 扩端"),
    ("05", "项目结构与优点", "apps、packages 每层功能；跟不跟趋势"),
    ("06", "如何延展", "App / iPad：加 Adapter，不重写 DSL 和控制台"),
]
for i, (n, title, desc) in enumerate(items):
    col, row = i % 2, i // 2
    x = Inches(0.5) + Inches(col * 6.4)
    y = Inches(1.5) + Inches(row * 1.75)
    round_box(s, x, y, Inches(6.1), Inches(1.55), WHITE, LINE)
    box(s, x, y, Inches(0.1), Inches(1.55), TEAL)
    tb(s, x + Inches(0.35), y + Inches(0.28), Inches(1.1), Inches(0.5), n, 24, True, TEAL)
    tb(s, x + Inches(1.5), y + Inches(0.28), Inches(4.3), Inches(0.45), title, 18, True, NAVY)
    tb(s, x + Inches(1.5), y + Inches(0.78), Inches(4.3), Inches(0.5), desc, 13, False, MUTED)


# ========== 3 Problem ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "业务测试今天真正卡住的地方", "不是缺工具，是缺「人和 AI 都能用」的统一测试基础设施")
footer(s, 3)
problems = [
    ("端各一套脚本", "小程序一套、Web 一套、接口再一套。跨端业务（下单→后台核验）要拼两套框架，变量与证据对不齐。"),
    ("AI 直接去点页面", "Agent 若直连 Playwright / 开发者工具，用例不可复现、不可校验、不可进 CI，也难做团队资产。"),
    ("用例散落、难共享", "脚本在个人仓库或 IDE 里；没有统一语言、没有元数据索引、Git 与执行结果脱节。"),
    ("业务能力被写进工具", "Login / 下单 / 支付做成「平台 Action」后，工具绑死业务，换项目就要改核心。"),
]
for i, (t, b) in enumerate(problems):
    x = Inches(0.45) + Inches((i % 4) * 3.2)
    round_box(s, x, Inches(1.55), Inches(3.05), Inches(4.9), WHITE, LINE)
    box(s, x, Inches(1.55), Inches(3.05), Inches(0.1), GOLD if i < 2 else TEAL)
    tb(s, x + Inches(0.18), Inches(1.85), Inches(2.7), Inches(1.1), t, 18, True, NAVY)
    tb(s, x + Inches(0.18), Inches(3.1), Inches(2.7), Inches(2.9), b, 13, False, MUTED)


# ========== 4 Positioning ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "一句话定位", "TestPilot 定义「怎么描述 / 校验 / 执行」；业务项目定义「测什么、怎么算通过」")
footer(s, 4)
round_box(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(1.55), NAVY)
tb(s, Inches(0.8), Inches(1.7), Inches(11.7), Inches(1.15), "TestPilot 是工具，不是业务测试项目。\n它让 AI 与工程师用同一套 DSL 写用例，用同一套引擎跨端执行，并把证据与结果沉淀为可审计资产。", 16, False, WHITE)

left = round_box(s, Inches(0.5), Inches(3.3), Inches(6.0), Inches(3.5), WHITE, LINE)
tb(s, Inches(0.75), Inches(3.5), Inches(5.5), Inches(0.4), "TestPilot 负责", 16, True, TEAL)
for i, line in enumerate(["测试描述语言（DSL）", "校验（schema / locator / adapter 能力）", "执行引擎与 Adapter 抽象", "取证与报告", "AI Skill（怎么用这套工具）", "控制面：CLI / Server / Web / Agent"]):
    tb(s, Inches(0.85), Inches(3.95) + Inches(i * 0.4), Inches(5.4), Inches(0.38), "▸  " + line, 13, False, INK)

right = round_box(s, Inches(6.8), Inches(3.3), Inches(6.0), Inches(3.5), WHITE, LINE)
tb(s, Inches(7.05), Inches(3.5), Inches(5.5), Inches(0.4), "业务项目负责", 16, True, GOLD)
for i, line in enumerate(["测哪条业务、哪个页面、什么预期", "Case 文件：tests/e2e/cases/*.yaml", "testpilot.yaml 与环境变量", "账号、域名、Workspace 组合", "Git 作为 Case 的 Source of Truth", "CI 里 validate / sync / run"]):
    tb(s, Inches(7.15), Inches(3.95) + Inches(i * 0.4), Inches(5.4), Inches(0.38), "▸  " + line, 13, False, INK)


# ========== 5 Principles ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "五条设计原则（后面所有结构都服从这五条）", "原则比框架重要：换 Playwright 版本可以，换原则会伤整条链路")
footer(s, 5)
principles = [
    ("01", "工具 ≠ 业务", "仓库里不放 login/order/payment 用例。业务 Case 永远在业务项目。"),
    ("02", "Skill 不是执行器", "AI 禁止直连 Playwright / WeChatIDE。固定：Skill → CLI/SDK → DSL → Engine → Adapter。"),
    ("03", "Case 是项目资产", "YAML 进 Git。Server 只存索引与 Run，不是第二套 Git。"),
    ("04", "共用模型，不共用驱动", "Web / 小程序 / API 共享 DSL、Context、Event、Artifact、Report；执行实现各自 Adapter。"),
    ("05", "Evidence-first", "AI 不能凭空编定位器。需求 → 代码/页面 → 证据 → 步骤 → Case。"),
]
for i, (n, t, b) in enumerate(principles):
    y = Inches(1.42) + Inches(i * 1.1)
    round_box(s, Inches(0.5), y, Inches(12.3), Inches(1.0), WHITE, LINE)
    round_box(s, Inches(0.7), y + Inches(0.22), Inches(0.85), Inches(0.55), TEAL)
    tb(s, Inches(0.7), y + Inches(0.22), Inches(0.85), Inches(0.55), n, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, Inches(1.75), y + Inches(0.12), Inches(10.7), Inches(0.38), t, 16, True, NAVY)
    tb(s, Inches(1.75), y + Inches(0.5), Inches(10.7), Inches(0.4), b, 13, False, MUTED)


# ========== 6 Architecture ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "核心架构：一条链，两端分叉", "人和 AI 走同一入口；真正碰页面的只有 Adapter")
footer(s, 6)

layers = [
    (0.5, "AI / 同事", "Cursor · Claude · 控制台"),
    (2.55, "Skill", "规范：怎么写 Case、禁直连驱动"),
    (4.6, "CLI / SDK", "csspilot · @testpilot/sdk"),
    (6.65, "DSL + Engine", "YAML 用例 · 调度 · 上下文 · 事件"),
]
for x, title, sub in layers:
    round_box(s, Inches(x), Inches(1.5), Inches(1.9), Inches(1.45), NAVY)
    tb(s, Inches(x), Inches(1.62), Inches(1.9), Inches(0.55), title, 14, True, WHITE, PP_ALIGN.CENTER)
    tb(s, Inches(x + 0.08), Inches(2.2), Inches(1.74), Inches(0.6), sub, 11, False, RGBColor(0x9E, 0xD4, 0xD6), PP_ALIGN.CENTER)
    if x < 6:
        tb(s, Inches(x + 1.85), Inches(1.9), Inches(0.28), Inches(0.5), "→", 18, True, TEAL, PP_ALIGN.CENTER)

# adapter fork
round_box(s, Inches(0.5), Inches(3.25), Inches(12.3), Inches(3.5), WHITE, LINE)
tb(s, Inches(0.75), Inches(3.4), Inches(12), Inches(0.35), "Adapter Resolver：按步骤 target 选执行器（同一条 Case 可切换端）", 14, True, NAVY)

ads = [
    ("web", "Playwright", "浏览器 / 后台"),
    ("miniapp", "WeChatIDE", "微信开发者工具"),
    ("api", "HTTP Adapter", "接口断言与抽取"),
]
for i, (tg, impl, note) in enumerate(ads):
    x = Inches(0.85) + Inches(i * 4.05)
    round_box(s, x, Inches(3.9), Inches(3.8), Inches(2.5), SOFT, TEAL)
    tb(s, x + Inches(0.2), Inches(4.05), Inches(3.4), Inches(0.4), f"target: {tg}", 14, True, TEAL)
    tb(s, x + Inches(0.2), Inches(4.5), Inches(3.4), Inches(0.45), impl, 20, True, NAVY)
    tb(s, x + Inches(0.2), Inches(5.1), Inches(3.4), Inches(0.9), note + "\n实现 TestAdapter 接口即可接入", 13, False, MUTED)


# ========== 7 Three nouns ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "先记住三个词：DSL、引擎、Adapter", "类比：菜谱 / 后厨流程 / 具体灶具。换灶具不必重写菜谱。")
footer(s, 7)
triples = [
    ("DSL", "测试说明书", "用 YAML 写：在哪个端、点什么、期望看到什么。人和 AI 都只写这一层。文件在业务项目 tests/e2e/cases/。"),
    ("Execution Engine", "执行调度", "按步骤读说明书：校验、解析 ${变量}、开会话、失败即停、记事件和产物。不关心按钮在浏览器还是小程序里。"),
    ("Adapter", "端上的手", "把「click .submit-btn」翻译成 Playwright 点击、或微信开发者工具点击、或以后的 Appium 点击。"),
]
for i, (n, role, b) in enumerate(triples):
    x = Inches(0.4) + Inches(i * 4.3)
    round_box(s, x, Inches(1.45), Inches(4.1), Inches(4.0), WHITE, LINE)
    box(s, x, Inches(1.45), Inches(4.1), Inches(0.9), NAVY)
    tb(s, x, Inches(1.5), Inches(4.1), Inches(0.45), n, 18, True, WHITE, PP_ALIGN.CENTER)
    tb(s, x, Inches(1.92), Inches(4.1), Inches(0.35), role, 13, False, TEAL, PP_ALIGN.CENTER)
    tb(s, x + Inches(0.2), Inches(2.55), Inches(3.7), Inches(2.7), b, 14, False, MUTED)
round_box(s, Inches(0.4), Inches(5.65), Inches(12.5), Inches(1.2), NAVY)
tb(s, Inches(0.65), Inches(5.85), Inches(12.1), Inches(0.85), "同事写/审的是 DSL。平台组维护引擎和 Adapter。AI 被规定：只能生成 DSL，禁止直接操作 Playwright / 微信开发者工具。", 15, False, WHITE)


# ========== FLOW 1 端到端主流程 ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "流程图 ①　从一条需求到一份报告", "四条泳道：谁在做 → 用什么工具 → 落到哪个端 → 产出什么")
footer(s)

lane_titles = ["人 / AI", "工具链", "端（驱动）", "产出"]
lane_y = [Inches(1.35), Inches(2.75), Inches(4.15), Inches(5.55)]
for i, name in enumerate(lane_titles):
    round_box(s, Inches(0.4), lane_y[i], Inches(1.5), Inches(1.15), NAVY)
    tb(s, Inches(0.4), lane_y[i], Inches(1.5), Inches(1.15), name, 13, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    box(s, Inches(2.0), lane_y[i], Inches(10.9), Inches(1.15), WHITE)

# lane 1
l1 = [(2.15, "需求 / 测试点"), (5.0, "读 Skill 规范\n（怎么写、先取证）"), (7.85, "写 Case YAML"), (10.7, "人工评审")]
for x, t in l1:
    node(s, Inches(x), lane_y[0] + Inches(0.18), Inches(2.3), Inches(0.8), t, WHITE, NAVY, 12)
for x in (4.55, 7.4, 10.25):
    ar_r(s, Inches(x), lane_y[0] + Inches(0.49))

# lane 2
l2 = [(2.15, "csspilot validate\n语法 / 定位 / 变量"), (5.0, "csspilot run\n（经 SDK）"), (7.85, "Execution Engine\n解析变量 · 开会话"), (10.7, "Adapter Resolver\n按 target 选驱动")]
for x, t in l2:
    node(s, Inches(x), lane_y[1] + Inches(0.18), Inches(2.3), Inches(0.8), t, SOFT, NAVY, 11)
for x in (4.55, 7.4, 10.25):
    ar_r(s, Inches(x), lane_y[1] + Inches(0.49))

# lane 3
l3 = [(2.15, "Playwright\ntarget: web"), (5.0, "WeChatIDE\ntarget: miniapp"), (7.85, "HTTP\ntarget: api"), (10.7, "（未来）App / iPad")]
for i, (x, t) in enumerate(l3):
    node(s, Inches(x), lane_y[2] + Inches(0.18), Inches(2.3), Inches(0.8), t, WHITE, NAVY if i < 3 else MUTED, 12, TEAL if i < 3 else LINE)

# lane 4
l4 = [(2.15, "Evidence\n截图 / 日志 / Trace"), (5.0, "Report\nJSON + HTML"), (7.85, "控制台 / CI\n团队可见"), (10.7, "失败归因 → 修用例")]
for x, t in l4:
    node(s, Inches(x), lane_y[3] + Inches(0.18), Inches(2.3), Inches(0.8), t, WHITE, NAVY, 12, TEAL)
for x in (4.55, 7.4, 10.25):
    ar_r(s, Inches(x), lane_y[3] + Inches(0.49))

# 跨泳道连接
for i in range(3):
    ar_d(s, Inches(12.45), lane_y[i] + Inches(1.0), Inches(0.18), Inches(0.28))


# ========== 8 How it works ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "一次运行实际发生了什么", "CLI 与 Server 都不直连引擎细节，一律走 SDK")
footer(s)
steps = [
    ("1", "Load", "读取 tests/e2e/cases\n解析 YAML"),
    ("2", "Validate", "schema / action\nlocator / 变量 / 端能力"),
    ("3", "Engine", "逐步执行\n维护 Context 与会话"),
    ("4", "Adapter", "把原语翻译成\n真实点击 / 请求"),
    ("5", "Evidence", "截图 · 日志\nvideo / trace"),
    ("6", "Report", "JSON + HTML\nStep 级结果"),
]
for i, (n, t, b) in enumerate(steps):
    x = Inches(0.4) + Inches(i * 2.15)
    round_box(s, x, Inches(1.55), Inches(2.0), Inches(2.7), WHITE, LINE)
    round_box(s, x + Inches(0.7), Inches(1.75), Inches(0.55), Inches(0.55), TEAL)
    tb(s, x + Inches(0.7), Inches(1.75), Inches(0.55), Inches(0.55), n, 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.1), Inches(2.4), Inches(1.8), Inches(0.4), t, 16, True, NAVY, PP_ALIGN.CENTER)
    tb(s, x + Inches(0.1), Inches(2.85), Inches(1.8), Inches(1.15), b, 12, False, MUTED, PP_ALIGN.CENTER)
    if i < 5:
        tb(s, x + Inches(1.88), Inches(2.5), Inches(0.3), Inches(0.4), "→", 16, True, TEAL)

round_box(s, Inches(0.4), Inches(4.5), Inches(12.5), Inches(2.35), WHITE, LINE)
tb(s, Inches(0.65), Inches(4.65), Inches(12), Inches(0.4), "关键约束", 15, True, NAVY)
cols = [
    ("Context 跨端传值", "extract 得到的 orderId 可在后续 web/api 步骤用 ${orderId}"),
    ("失败即停后续", "未执行步骤标记 skipped；取消信号可中止剩余用例"),
    ("事件可观测", "NDJSON 事件流：Web / CLI 都能看 Step 进度"),
    ("产物不进库", "截图/视频留在业务项目 .testpilot/artifacts/"),
]
for i, (t, b) in enumerate(cols):
    x = Inches(0.65) + Inches(i * 3.1)
    tb(s, x, Inches(5.15), Inches(2.95), Inches(0.4), t, 13, True, TEAL)
    tb(s, x, Inches(5.55), Inches(2.95), Inches(1.1), b, 12, False, MUTED)


# ========== FLOW 2 单次 Run 执行流程 ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "流程图 ②　单次 Run 的执行判断", "校验不过就不执行；某步失败即停，后续步骤记为 skipped，但证据一定留下")
footer(s)

y1 = Inches(1.45)
node(s, Inches(0.45), y1, Inches(1.9), Inches(0.7), "加载 Case\nYAML → 对象", WHITE, NAVY, 12)
ar_r(s, Inches(2.45), y1 + Inches(0.26))
node(s, Inches(2.85), y1, Inches(1.9), Inches(0.7), "Validate\nschema / locator", WHITE, NAVY, 12)
ar_r(s, Inches(4.85), y1 + Inches(0.26))
diamond(s, Inches(5.25), Inches(1.28), Inches(2.0), Inches(1.05), "校验通过?")
ar_r(s, Inches(7.35), y1 + Inches(0.26))
ar_label(s, Inches(7.3), Inches(1.2), "否", GOLD)
node(s, Inches(7.75), y1, Inches(2.3), Inches(0.7), "报错并退出\nexit code 1", WHITE, RGBColor(0xC0, 0x4B, 0x3A), 12, RGBColor(0xC0, 0x4B, 0x3A))
tb(s, Inches(10.2), y1, Inches(2.7), Inches(0.7), "validate 在 CI 里也单独跑，\n坏用例进不了主干。", 11, False, MUTED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

ar_d(s, Inches(6.16), Inches(2.4), Inches(0.18), Inches(0.28))
ar_label(s, Inches(6.3), Inches(2.42), "是", TEAL)

y2 = Inches(2.78)
node(s, Inches(2.85), y2, Inches(2.4), Inches(0.7), "创建 ExecutionContext\n变量表 / 运行目录", SOFT, NAVY, 11)
ar_r(s, Inches(5.35), y2 + Inches(0.26))
node(s, Inches(5.75), y2, Inches(2.4), Inches(0.7), "按 target 取会话\nAdapter Resolver", SOFT, NAVY, 11)
ar_r(s, Inches(8.25), y2 + Inches(0.26))
node(s, Inches(8.65), y2, Inches(2.4), Inches(0.7), "执行一个 Step\n解析 ${变量} 后下发", NAVY, WHITE, 11, None)

ar_d(s, Inches(9.76), Inches(3.55), Inches(0.18), Inches(0.28))

y3 = Inches(3.95)
diamond(s, Inches(8.85), y3, Inches(2.0), Inches(1.05), "该步成功?")
ar_l(s, Inches(8.35), y3 + Inches(0.44))
ar_label(s, Inches(8.25), Inches(3.72), "否", GOLD)
node(s, Inches(5.3), y3 + Inches(0.17), Inches(2.9), Inches(0.7), "截图取证 + 记录错误\n该步 failed", WHITE, RGBColor(0xC0, 0x4B, 0x3A), 12, RGBColor(0xC0, 0x4B, 0x3A))
ar_l(s, Inches(4.9), y3 + Inches(0.44))
node(s, Inches(2.0), y3 + Inches(0.17), Inches(2.8), Inches(0.7), "剩余 Step → skipped\nCase 判为 failed", WHITE, NAVY, 12)

ar_d(s, Inches(9.76), y3 + Inches(1.1), Inches(0.18), Inches(0.28))
ar_label(s, Inches(9.9), y3 + Inches(1.12), "是", TEAL)

y4 = Inches(5.5)
diamond(s, Inches(8.85), y4, Inches(2.0), Inches(1.05), "还有下一步?")
tb(s, Inches(11.1), y4 + Inches(0.15), Inches(2.0), Inches(0.75), "是 → 回到「执行一个 Step」\n（循环）", 11, True, TEAL, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
ar_l(s, Inches(8.35), y4 + Inches(0.44))
ar_label(s, Inches(8.2), y4 - Inches(0.22), "否", GOLD)
node(s, Inches(5.3), y4 + Inches(0.17), Inches(2.9), Inches(0.7), "汇总 CaseResult\n事件流 NDJSON", SOFT, NAVY, 12)
ar_l(s, Inches(4.9), y4 + Inches(0.44))
node(s, Inches(2.0), y4 + Inches(0.17), Inches(2.8), Inches(0.7), "Report\nJSON + HTML", WHITE, NAVY, 12, TEAL)


# ========== 9 DSL what ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "DSL 是什么", "Domain Specific Language：专门用来写「业务测试步骤」的小语言，不是通用编程语言")
footer(s, 9)
round_box(s, Inches(0.4), Inches(1.4), Inches(12.5), Inches(1.45), NAVY)
tb(s, Inches(0.65), Inches(1.55), Inches(12.1), Inches(1.15), "一句话：DSL 就是一份 YAML 测试说明书。\n写的是「在哪个端、做哪个动作、点哪个控件、期望看到什么」，不写 Playwright 代码，也不写小程序 automator 代码。", 16, False, WHITE)
pairs = [
    ("它是", "结构化的 Case 文件（id / name / steps）。代码里用 zod 校验对不对。包：packages/dsl。"),
    ("它不是", "JavaScript 测试脚本，也不是「登录」「下单」这种业务按钮。那些是用原语拼出来的业务 Case。"),
    ("谁来写", "测试同学手写，或 AI 按 Skill 生成。生成后必须 csspilot validate，才能 run。"),
    ("放在哪", "业务项目 tests/e2e/cases/*.yaml。TestPilot 仓库只放示例，不拥有业务用例。"),
]
for i, (t, b) in enumerate(pairs):
    col, row = i % 2, i // 2
    x = Inches(0.4) + Inches(col * 6.45)
    y = Inches(3.05) + Inches(row * 1.9)
    round_box(s, x, y, Inches(6.25), Inches(1.75), WHITE, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.2), Inches(5.8), Inches(0.4), t, 16, True, TEAL)
    tb(s, x + Inches(0.25), y + Inches(0.65), Inches(5.8), Inches(0.9), b, 13, False, MUTED)


# ========== 10 DSL anatomy ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "一条 DSL Case 由什么组成", "每一步只有三件关键信息：在哪一端（target）· 做什么（action）· 点哪里/期望什么")
footer(s, 10)
parts = [
    ("Case 头", "id / name / tags\n可选 workspace、environment、accountRef\nCase 不写死 URL 和密码"),
    ("target 执行端", "web 浏览器/后台\nminiapp 微信小程序\napi HTTP 接口\n一步一个端，一条 Case 可换端"),
    ("action 原语", "十个稳定动作：\nlaunch navigate click input select\nwait assert extract screenshot request"),
    ("定位与数据", "locator：css 或 text\nextract 写入变量 ${orderId}\n后续步骤可跨端引用"),
]
for i, (t, b) in enumerate(parts):
    x = Inches(0.4) + Inches(i * 3.2)
    round_box(s, x, Inches(1.45), Inches(3.05), Inches(3.55), WHITE, LINE)
    box(s, x, Inches(1.45), Inches(3.05), Inches(0.08), TEAL)
    tb(s, x + Inches(0.15), Inches(1.65), Inches(2.75), Inches(0.7), t, 15, True, NAVY)
    tb(s, x + Inches(0.15), Inches(2.4), Inches(2.75), Inches(2.35), b, 13, False, MUTED)
round_box(s, Inches(0.4), Inches(5.2), Inches(12.5), Inches(1.65), WHITE, LINE)
tb(s, Inches(0.65), Inches(5.35), Inches(12.1), Inches(0.4), "刻意不含的东西", 15, True, GOLD)
tb(s, Inches(0.65), Inches(5.8), Inches(12.1), Inches(0.85), "没有 login / order / payment 这种业务 Action。业务同学用 click+input+assert 组合出「登录」；换一个业务项目，DSL 原语不用改，改的是 YAML 内容。", 14, False, MUTED)


# ========== 11 DSL example ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "读一段真实 DSL：跨端黄金路径", "同一文件里换 target，变量跨端传递——这就是「共用测试模型」")
footer(s, 11)

round_box(s, Inches(0.4), Inches(1.45), Inches(7.4), Inches(5.4), WHITE, LINE)
tb(s, Inches(0.6), Inches(1.58), Inches(7), Inches(0.35), "黄金路径示例（小程序下单 → Web 核验）", 14, True, NAVY)
dsl = """id: order-create
name: 用户创建订单
steps:
  - target: miniapp
    action: launch
  - target: miniapp
    action: click
    locator: { css: ".submit-btn" }
  - target: miniapp
    action: extract
    locator: { css: ".order-id" }
    variable: orderId
  - target: web
    action: assert
    locator: { css: ".order-row-${orderId}" }
    expected: "${orderId}" """
tb(s, Inches(0.65), Inches(2.0), Inches(7.0), Inches(4.6), dsl, 13, False, INK)

right_items = [
    ("执行端 target", "web  ·  miniapp  ·  api\n扩端时在此增加枚举"),
    ("十个 Action", "launch navigate click input\nselect wait assert extract\nscreenshot request"),
    ("变量与环境", "${VAR} 模板；environment /\nworkspace / accountRef 引用环境，不写死 URL 与密码"),
]
for i, (t, b) in enumerate(right_items):
    y = Inches(1.45) + Inches(i * 1.85)
    round_box(s, Inches(8.0), y, Inches(4.9), Inches(1.7), WHITE, LINE)
    tb(s, Inches(8.2), y + Inches(0.15), Inches(4.5), Inches(0.4), t, 14, True, TEAL)
    tb(s, Inches(8.2), y + Inches(0.55), Inches(4.5), Inches(1.0), b, 13, False, MUTED)


# ========== 12 Adapter what ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "Adapter 是什么", "适配器：引擎只认识统一接口 TestAdapter；每个端自己实现「这句话在我这边怎么执行」")
footer(s, 12)
round_box(s, Inches(0.4), Inches(1.4), Inches(12.5), Inches(1.4), NAVY)
tb(s, Inches(0.65), Inches(1.55), Inches(12.1), Inches(1.1), "一句话：Adapter 是「翻译官 + 机械手」。\nDSL 说 click，Web Adapter 调用 Playwright；小程序 Adapter 调用微信开发者工具；API Adapter 发 HTTP。引擎从不 import 这些工具。", 16, False, WHITE)
cols = [
    ("像电源转换头", "墙上插座规格统一（TestAdapter 方法）。各国插头不同（浏览器 / 小程序 / App）。换插头，不换家里的电器。"),
    ("引擎依赖抽象", "packages/adapter-core 定义 launch/click/assert…\nexecution-engine 只调这个接口。"),
    ("谁来实现", "packages/adapter-playwright\npackages/adapter-wechatide\npackages/adapter-api\n以后 adapter-app / adapter-ipad"),
]
for i, (t, b) in enumerate(cols):
    x = Inches(0.4) + Inches(i * 4.3)
    round_box(s, x, Inches(3.05), Inches(4.1), Inches(3.75), WHITE, LINE)
    tb(s, x + Inches(0.2), Inches(3.25), Inches(3.7), Inches(0.7), t, 16, True, TEAL)
    tb(s, x + Inches(0.2), Inches(4.05), Inches(3.7), Inches(2.5), b, 14, False, MUTED)


# ========== 13 Three adapters ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "现在三套 Adapter 各自干什么", "同一条 Case 可以交替使用它们；Resolver 按步骤的 target 选中对应会话")
footer(s, 13)
ads2 = [
    ("adapter-playwright", "target: web", "驱动 Chromium 等浏览器。打开后台、点击、输入、断言文本、截图、可选 video/trace。对应业务：管理端、H5、PC Web。"),
    ("adapter-wechatide", "target: miniapp", "驱动微信开发者工具（miniprogram-automator）。launch 小程序项目、navigate 页面、css 定位点击。对应业务：微信小程序。"),
    ("adapter-api", "target: api", "发 HTTP（request），对响应 assert / 按 JSON path extract。可选能力，UI Adapter 不必实现。对应：直接验接口或串数据。"),
]
for i, (n, tg, b) in enumerate(ads2):
    y = Inches(1.42) + Inches(i * 1.75)
    round_box(s, Inches(0.4), y, Inches(12.5), Inches(1.6), WHITE, LINE)
    box(s, Inches(0.4), y, Inches(0.1), Inches(1.6), TEAL)
    tb(s, Inches(0.7), y + Inches(0.18), Inches(7.5), Inches(0.4), n, 16, True, NAVY)
    tb(s, Inches(8.4), y + Inches(0.18), Inches(4.2), Inches(0.4), tg, 14, True, TEAL)
    tb(s, Inches(0.7), y + Inches(0.65), Inches(11.9), Inches(0.8), b, 14, False, MUTED)


# ========== 14 Adapter contract ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "为什么「加一端」不必重写测试体系", "引擎只认 TestAdapter；Playwright / 微信工具 / HTTP 都是插件")
footer(s, 14)
tb(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.4), "接口契约（packages/adapter-core）：launch / navigate / click / input / select / wait / assert / extract / screenshot / close", 13, False, MUTED)

iface = [
    ("引擎看到的", "StepTarget + Action\n会话、超时、变量已解析"),
    ("Adapter 翻译的", "CSS/text → 真实控件\n坐标、WebView、HTTP…"),
    ("可选能力", "startEvidence / video\nrequest / extractResponse"),
]
for i, (t, b) in enumerate(iface):
    x = Inches(0.5) + Inches(i * 4.2)
    round_box(s, x, Inches(1.95), Inches(4.0), Inches(1.7), WHITE, LINE)
    tb(s, x + Inches(0.2), Inches(2.1), Inches(3.6), Inches(0.4), t, 15, True, NAVY)
    tb(s, x + Inches(0.2), Inches(2.55), Inches(3.6), Inches(0.9), b, 13, False, MUTED)

round_box(s, Inches(0.5), Inches(3.9), Inches(12.3), Inches(2.9), NAVY)
tb(s, Inches(0.8), Inches(4.1), Inches(11.8), Inches(0.4), "对领导的含义", 16, True, TEAL)
tb(
    s,
    Inches(0.8),
    Inches(4.6),
    Inches(11.8),
    Inches(1.9),
    "今天已验证：同一条 Case 里 miniapp 与 web 来回切换，订单号跨端传递。\n以后加 App / iPad，是「再实现一个 Adapter + 扩展 target 枚举」，而不是再造一套用例语言、再造一套报告、再教 AI 另一套 Skill。\n存量 YAML 业务资产可以继续用；新端只补「怎么点到那个按钮」这一层。",
    15,
    False,
    WHITE,
)


# ========== 10 Control plane ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "控制面：本地可跑，团队可看", "增量层不推翻 Core：Web / Server / Agent 全部经 SDK 访问引擎")
footer(s, 15)
planes = [
    ("CLI  csspilot", "init / validate / list / run\nreport / doctor / update\nci init / sync-metadata\n个人与 CI 的主入口"),
    ("SDK", "唯一编程接口\n禁止 Server/Agent 直连\nexecution-engine 或驱动\n保证行为一致"),
    ("Server", "Fastify + SQLite\n项目 / Case 索引 / Run\nAgent Job / Workspace\n静态托管 Web 控制台"),
    ("Web 控制台", "项目概览 · Case Index\nRun 详情（Step + 截图）\n同步状态 · 报告下载\n给测试/研发/领导看结果"),
    ("Agent", "Planner：规划并写 YAML\nAnalysis：失败归因建议\n落盘在业务项目 root\n可见性仍走 git push"),
]
for i, (t, b) in enumerate(planes):
    x = Inches(0.35) + Inches(i * 2.58)
    round_box(s, x, Inches(1.5), Inches(2.45), Inches(5.3), WHITE, LINE)
    box(s, x, Inches(1.5), Inches(2.45), Inches(0.1), TEAL)
    tb(s, x + Inches(0.12), Inches(1.75), Inches(2.2), Inches(0.9), t, 15, True, NAVY)
    tb(s, x + Inches(0.12), Inches(2.75), Inches(2.2), Inches(3.7), b, 13, False, MUTED)


# ========== FLOW 3 AI 生成与失败闭环 ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "流程图 ③　AI 参与的闭环：生成 → 校验 → 执行 → 归因 → 修正", "AI 只在两处介入：写 Case 和解释失败。中间的执行与判定全部由确定性的引擎完成。")
footer(s)

top = Inches(1.45)
seq = [
    (0.45, "提出测试需求\n（人）", WHITE, NAVY),
    (2.75, "Planner Agent\n读代码/页面取证", SOFT, NAVY),
    (5.05, "生成 Case YAML\n落在业务项目", WHITE, NAVY),
    (7.35, "自动 validate", WHITE, NAVY),
    (9.65, "可选：立即 run", NAVY, WHITE),
]
for x, t, f, c in seq:
    node(s, Inches(x), top, Inches(2.1), Inches(0.85), t, f, c, 12, None if f == NAVY else LINE)
for x in (2.35, 4.65, 6.95, 9.25):
    ar_r(s, Inches(x), top + Inches(0.33))
ar_r(s, Inches(11.85), top + Inches(0.33), Inches(0.35), Inches(0.2))
tb(s, Inches(12.25), top, Inches(1.0), Inches(0.85), "结果", 12, True, NAVY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

ar_d(s, Inches(10.6), Inches(2.4), Inches(0.18), Inches(0.3))
mid = Inches(2.85)
diamond(s, Inches(9.65), mid, Inches(2.1), Inches(1.1), "全部通过?")
ar_r(s, Inches(11.8), mid + Inches(0.46))
ar_label(s, Inches(11.7), mid - Inches(0.22), "是", TEAL)
node(s, Inches(12.2), mid + Inches(0.2), Inches(0.95), Inches(0.7), "合入\nPR", WHITE, NAVY, 11, TEAL)

ar_l(s, Inches(9.15), mid + Inches(0.46))
ar_label(s, Inches(9.0), mid - Inches(0.22), "否", GOLD)
node(s, Inches(6.55), mid + Inches(0.2), Inches(2.5), Inches(0.7), "Evidence：截图 / 日志\nStep 级失败点", WHITE, RGBColor(0xC0, 0x4B, 0x3A), 11, RGBColor(0xC0, 0x4B, 0x3A))
ar_l(s, Inches(6.15), mid + Inches(0.46))
node(s, Inches(3.55), mid + Inches(0.2), Inches(2.5), Inches(0.7), "Analysis Agent\n归因并给建议", SOFT, NAVY, 11)
ar_l(s, Inches(3.15), mid + Inches(0.46))
diamond(s, Inches(0.7), mid - Inches(0.05), Inches(2.3), Inches(1.2), "是用例问题\n还是真 Bug?")

bot = Inches(4.65)
ar_d(s, Inches(1.76), Inches(4.25), Inches(0.18), Inches(0.3))
node(s, Inches(0.45), bot, Inches(2.8), Inches(0.9), "用例问题 → 改 locator / 步骤\n重新 validate & run", WHITE, NAVY, 12, TEAL)
ar_r(s, Inches(3.35), bot + Inches(0.35))
node(s, Inches(3.75), bot, Inches(2.8), Inches(0.9), "真 Bug → 提缺陷单\n附截图 / Trace 证据", WHITE, NAVY, 12, GOLD)

round_box(s, Inches(7.0), Inches(4.55), Inches(5.9), Inches(2.1), NAVY)
tb(s, Inches(7.25), Inches(4.72), Inches(5.4), Inches(0.35), "为什么这样切分", 14, True, TEAL)
tb(s, Inches(7.25), Inches(5.12), Inches(5.4), Inches(1.4), "AI 擅长「写」和「解释」，不擅长「保证」。\n所以判定通过与否的是引擎和断言，不是模型；\nAI 的产出必须先过 validate，才允许执行。\n这样 AI 参与量可以放大，风险不会同比放大。", 13, False, WHITE)


# ========== 11 Git ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "Git 是 Case 的真相源，Server 不是第二套 Git", "分支即环境；Workspace 绑定多系统组合；历史 Run 可还原当时快照")
footer(s, 16)
flow = [
    ("业务仓库", "YAML Case\n在 Git 里评审"),
    ("git push", "走正常代码\n评审与保护分支"),
    ("CI", "validate +\nsync-metadata"),
    ("Server 索引", "快照同步\n消失则标记 deleted"),
    ("控制台可见", "团队看 Index\n与历史 Run"),
]
for i, (t, b) in enumerate(flow):
    x = Inches(0.4) + Inches(i * 2.55)
    round_box(s, x, Inches(1.5), Inches(2.35), Inches(1.9), NAVY if i in (2, 3) else WHITE, TEAL if i not in (2, 3) else None)
    c = WHITE if i in (2, 3) else NAVY
    m = RGBColor(0x9E, 0xD4, 0xD6) if i in (2, 3) else MUTED
    tb(s, x + Inches(0.1), Inches(1.65), Inches(2.15), Inches(0.5), t, 14, True, c, PP_ALIGN.CENTER)
    tb(s, x + Inches(0.1), Inches(2.2), Inches(2.15), Inches(1.0), b, 12, False, m, PP_ALIGN.CENTER)
    if i < 4:
        tb(s, x + Inches(2.2), Inches(2.1), Inches(0.4), Inches(0.4), "→", 16, True, TEAL)

points = [
    ("分支即测试环境", "同一 Case 可同时存在于 main→TEST、release→STAGING。快照按分支作用域同步，互不影响。"),
    ("Workspace", "一次运行所需的多系统环境组合（小程序 + 后台 + API）。Case 只写 workspace 名称，不硬编码 URL/账号。"),
    ("Run 可追溯", "每次 Run 记录 branch / commit / workspace 运行时快照。环境后来改了，历史仍能还原。"),
    ("数据边界", "SQLite 只存元数据。大文件（截图、视频、trace）留在业务项目磁盘。testpilot.yaml 仍是配置真相源。"),
]
for i, (t, b) in enumerate(points):
    col, row = i % 2, i // 2
    x = Inches(0.4) + Inches(col * 6.45)
    y = Inches(3.7) + Inches(row * 1.55)
    round_box(s, x, y, Inches(6.25), Inches(1.4), WHITE, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.15), Inches(5.8), Inches(0.35), t, 14, True, TEAL)
    tb(s, x + Inches(0.25), y + Inches(0.55), Inches(5.8), Inches(0.7), b, 12, False, MUTED)


# ========== 17 Structure map ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "项目结构：先看「谁调用谁」", "上面是给人/AI 用的入口；中间是语言和引擎；下面才是真正碰端的 Adapter")
footer(s, 17)
layers2 = [
    ("入口层 apps + skills", "cli 命令行 · server API · web 控制台 · agent 规划/分析 · skills 教 AI 怎么用"),
    ("门面 SDK", "@testpilot/sdk：所有入口只能走这里访问 Core。禁止 Server/Agent 直接碰引擎或 Playwright。"),
    ("核心层 packages", "dsl 说明书规则 · execution-engine 调度 · core 配置类型 · evidence 取证 · reporter 报告"),
    ("驱动层 adapters", "adapter-core 接口；playwright / wechatide / api 三种实现。扩 App 只加这一层。"),
]
for i, (t, b) in enumerate(layers2):
    y = Inches(1.4) + Inches(i * 1.3)
    round_box(s, Inches(0.4), y, Inches(12.5), Inches(1.15), WHITE, LINE)
    round_box(s, Inches(0.55), y + Inches(0.28), Inches(0.55), Inches(0.55), TEAL)
    tb(s, Inches(0.55), y + Inches(0.28), Inches(0.55), Inches(0.55), str(i + 1), 16, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, Inches(1.3), y + Inches(0.15), Inches(11.3), Inches(0.4), t, 16, True, NAVY)
    tb(s, Inches(1.3), y + Inches(0.58), Inches(11.3), Inches(0.45), b, 13, False, MUTED)


# ========== 18 apps functions ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "apps/ 四个应用各自的功能", "这是「产品形态」：命令行、服务、页面、AI 进程。它们都不自己实现点击页面。")
footer(s, 18)
apps = [
    ("cli  （发布名 csspilot）", "给人/CI 用的命令。init 把测试基建接入业务仓；validate 检查 YAML；run 执行；report 出报告；doctor 体检；sync-metadata 把 Case 索引同步到 Server。"),
    ("server  Control Plane", "团队共享的 API：项目、Case 索引、Run、取消、产物、Agent Job、Workspace。SQLite 只存元数据。生产环境顺带托管 web 静态页。"),
    ("web  测试控制台", "给人看的界面：项目概览、Case 列表、某次 Run 的每一步结果、截图/日志/Trace。不执行测试，只展示 Server 的数据。"),
    ("agent  TestPilot AI", "Planner：按需求写 YAML 并 validate；Analysis：根据失败证据给建议。经 Skill → SDK 工作，YAML 写在业务项目磁盘，不直接写入 Server 索引。"),
]
for i, (t, b) in enumerate(apps):
    y = Inches(1.4) + Inches(i * 1.35)
    round_box(s, Inches(0.4), y, Inches(12.5), Inches(1.22), WHITE, LINE)
    tb(s, Inches(0.65), y + Inches(0.12), Inches(12.1), Inches(0.38), t, 15, True, TEAL)
    tb(s, Inches(0.65), y + Inches(0.52), Inches(12.1), Inches(0.6), b, 13, False, MUTED)


# ========== 19 packages functions ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "packages/ 每个包的功能", "这是「能力零件」。同事改业务 Case 不必动这里；平台扩端主要动 adapter-* 和 dsl 的 target 枚举。")
footer(s, 19)
pkgs = [
    ("core", "公共类型、错误码、testpilot.yaml 配置模型。大家说的「一次 Run / 一步结果」长什么样，在这里定。"),
    ("dsl", "Case 的语法：YAML 解析 + zod schema。决定哪些字段合法、target/action 有哪些。"),
    ("execution-engine", "真正跑步骤：解析变量、调 Resolver 拿会话、发 NDJSON 事件、收集结果、管理产物目录。"),
    ("adapter-core", "TestAdapter 接口 + 会话。所有端必须实现同一套方法，引擎才认。"),
    ("adapter-playwright / wechatide / api", "三个具体翻译官。分别把原语变成浏览器操作、小程序操作、HTTP 请求。"),
    ("sdk", "给 cli/server/agent 的统一编程入口：跑用例、取消、读报告。保证三条入口行为一致。"),
    ("evidence", "截图、视频、trace、日志怎么落盘。证据在业务项目 .testpilot/artifacts/，不进 Git。"),
    ("reporter", "把一次 Run 收成 JSON + HTML 报告，供本地打开或控制台展示。"),
]
for i, (t, b) in enumerate(pkgs):
    col, row = i % 2, i // 2
    x = Inches(0.35) + Inches(col * 6.5)
    y = Inches(1.38) + Inches(row * 1.35)
    round_box(s, x, y, Inches(6.3), Inches(1.25), WHITE, LINE)
    tb(s, x + Inches(0.2), y + Inches(0.1), Inches(5.95), Inches(0.38), t, 13, True, TEAL)
    tb(s, x + Inches(0.2), y + Inches(0.5), Inches(5.95), Inches(0.65), b, 12, False, MUTED)


# ========== 13 How we built ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "怎样做成的：先 Core，再控制面", "V0.1 明确先跑通跨端黄金 Case，再加 Dashboard / Agent，避免空中楼阁")
footer(s, 20)
left_ph = [
    ("Phase 1–2", "骨架 + CLI 六命令"),
    ("Phase 3–4", "DSL + Execution Engine"),
    ("Phase 5–6", "Playwright / WeChatIDE Adapter"),
    ("Phase 7", "跨端黄金 Case 真机验证"),
    ("Phase 8", "Skill 与实现对齐"),
]
right_ph = [
    ("Phase 9", "SDK，CLI 统一走 SDK；Run 可取消"),
    ("Phase 10", "Backend + Git 元数据快照同步"),
    ("Phase 11", "Web 控制台（Step 级证据）"),
    ("Phase 12", "Agent：写 Case + 失败分析"),
    ("发布", "csspilot npm；init 把 Skill 装进业务项目"),
]
tb(s, Inches(0.5), Inches(1.4), Inches(6), Inches(0.4), "Core 先行（必须可本地闭环）", 16, True, NAVY)
tb(s, Inches(6.9), Inches(1.4), Inches(6), Inches(0.4), "增量控制面（不推翻 Core）", 16, True, NAVY)
for i, (a, b) in enumerate(left_ph):
    y = Inches(1.9) + Inches(i * 0.95)
    round_box(s, Inches(0.5), y, Inches(6.1), Inches(0.85), WHITE, LINE)
    tb(s, Inches(0.7), y + Inches(0.22), Inches(1.8), Inches(0.4), a, 13, True, TEAL)
    tb(s, Inches(2.5), y + Inches(0.22), Inches(3.9), Inches(0.4), b, 14, False, INK)
for i, (a, b) in enumerate(right_ph):
    y = Inches(1.9) + Inches(i * 0.95)
    round_box(s, Inches(6.9), y, Inches(6.0), Inches(0.85), WHITE, LINE)
    tb(s, Inches(7.1), y + Inches(0.22), Inches(1.5), Inches(0.4), a, 13, True, GOLD)
    tb(s, Inches(8.7), y + Inches(0.22), Inches(4.0), Inches(0.4), b, 14, False, INK)


# ========== 14 Advantages ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "设计优点：给领导和协作方的清单", "每一条都能对应到代码边界，而不是口号")
footer(s, 21)
advs = [
    ("对业务团队", "一条 YAML 覆盖多端流程；Case 跟代码一起评审；不绑定某个人的 IDE 脚本。"),
    ("对 AI 研发", "Skill 规定「只能写 DSL、必须先取证」；生成物可 validate，能进 CI，而不是一次性点点点。"),
    ("对平台演进", "新端 = 新 Adapter。引擎、报告、控制台、Skill 工作流可复用。"),
    ("对质量治理", "Evidence-first + Step 级报告 + 取消与跳过语义清晰，失败可定位。"),
    ("对安全与合规", "凭据走环境变量 / Secret；HTTP 头脱敏；Case 里禁止写密码。"),
    ("对接入成本", "npx csspilot init 幂等接入；不覆盖已有测试；doctor 分组体检。"),
    ("对部署形态", "个人只跑 CLI 即可；团队加一个 Server 进程托管控制台，SQLite 起步。"),
    ("对长期所有权", "业务 Case 留在业务仓。换测试平台时，YAML 资产仍可读、可迁移。"),
]
for i, (t, b) in enumerate(advs):
    col, row = i % 4, i // 4
    x = Inches(0.35) + Inches(col * 3.25)
    y = Inches(1.45) + Inches(row * 2.7)
    round_box(s, x, y, Inches(3.1), Inches(2.5), WHITE, LINE)
    box(s, x, y, Inches(3.1), Inches(0.08), TEAL if row == 0 else GOLD)
    tb(s, x + Inches(0.15), y + Inches(0.25), Inches(2.8), Inches(0.7), t, 15, True, NAVY)
    tb(s, x + Inches(0.15), y + Inches(1.0), Inches(2.8), Inches(1.3), b, 12, False, MUTED)


# ========== 15 Trends ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "是否跟随趋势？跟随什么，刻意不跟什么", "对齐行业方向，但拒绝把 Demo 能力写进 Core")
footer(s, 22)
rows = [
    ("AI Native 测试", "跟随", "不是「用 AI 点页面」，而是给 AI 可校验的语言、规范与工具链（Skill + DSL + validate）。"),
    ("平台工程 / 内部开发者平台", "跟随", "测试能力产品化：init 接入、doctor、CI 模板、控制台，而不是每个组自建脚本。"),
    ("GitOps 式资产", "跟随", "Case 以 Git 为 SoT；push 后才进团队索引。与 IaC / 配置即代码同一治理逻辑。"),
    ("可插拔驱动（Adapter）", "跟随", "与浏览器自动化、Appium、多端 RPA 的分层一致：协议层稳定，驱动可换。"),
    ("全自动修 bug / RAG 知识库", "不跟（V0.1）", "未经验证的自动修复会污染仓库。先保证生成可跑、失败可解释。"),
    ("把业务封装成平台 Action", "明确反对", "Login/Order 进 Core 会让工具绑死业务。趋势是「领域用例在业务侧组合原语」。"),
]
tb(s, Inches(0.45), Inches(1.35), Inches(2.4), Inches(0.35), "议题", 12, True, MUTED)
tb(s, Inches(3.0), Inches(1.35), Inches(1.6), Inches(0.35), "态度", 12, True, MUTED)
tb(s, Inches(4.8), Inches(1.35), Inches(8), Inches(0.35), "说明", 12, True, MUTED)
for i, (a, b, c) in enumerate(rows):
    y = Inches(1.7) + Inches(i * 0.85)
    round_box(s, Inches(0.4), y, Inches(12.5), Inches(0.78), WHITE, LINE)
    tb(s, Inches(0.55), y + Inches(0.18), Inches(2.4), Inches(0.45), a, 13, True, NAVY)
    color = TEAL if b.startswith("跟随") else (GOLD if "不跟" in b else RGBColor(0xC0, 0x4B, 0x3A))
    tb(s, Inches(3.05), y + Inches(0.18), Inches(1.6), Inches(0.45), b, 13, True, color)
    tb(s, Inches(4.8), y + Inches(0.12), Inches(7.9), Inches(0.55), c, 12, False, MUTED)


# ========== 16 Current capabilities ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "V0.1 已经能做什么（基于仓库现状）", "不是规划图：跨端黄金路径与控制面均已落地")
footer(s, 23)
caps = [
    ("已通", "Web（Playwright）九个 UI Action + 取证"),
    ("已通", "小程序（WeChatIDE / automator）真机验证"),
    ("已通", "API target：request / 响应断言与抽取"),
    ("已通", "跨端变量：小程序 extract → Web assert"),
    ("已通", "CLI 全命令 + doctor / CI 模板 / npm 发布"),
    ("已通", "控制台：Case Index、Run、截图/Trace"),
    ("已通", "Agent 写 YAML + 失败分析（经 Skill/SDK）"),
    ("边界", "不做多用户权限、不做业务 Skill、不做自动修代码"),
]
for i, (tag, text) in enumerate(caps):
    col, row = i % 2, i // 2
    x = Inches(0.45) + Inches(col * 6.45)
    y = Inches(1.45) + Inches(row * 1.3)
    round_box(s, x, y, Inches(6.25), Inches(1.15), WHITE, LINE)
    bg = TEAL if tag == "已通" else GOLD
    round_box(s, x + Inches(0.2), y + Inches(0.32), Inches(1.05), Inches(0.5), bg)
    tb(s, x + Inches(0.2), y + Inches(0.32), Inches(1.05), Inches(0.5), tag, 13, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(1.45), y + Inches(0.3), Inches(4.55), Inches(0.55), text, 14, False, INK)


# ========== 17 Extend overview ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "以后如何延展到 App、iPad 和其它端", "稳定的是「测试模型」；变化的是「怎么驱动那块屏幕」")
footer(s, 24)

tb(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.4), "推荐节奏：先协议，再驱动，最后才改 Skill 示例。不要先买一套新的 App 测试平台另起炉灶。", 14, False, MUTED)

stages = [
    ("现在", TEAL, "web / miniapp / api\n同一 Case 可混排"),
    ("近", GOLD, "App（Android / iOS）\n实现 TestAdapter"),
    ("中", NAVY, "iPad / 折叠屏\n多窗口与尺寸策略"),
    ("远", MUTED, "桌面端 / 车机 / 鸿蒙\n仍是 target + Adapter"),
]
for i, (t, c, b) in enumerate(stages):
    x = Inches(0.45) + Inches(i * 3.2)
    round_box(s, x, Inches(1.95), Inches(3.0), Inches(2.35), WHITE, LINE)
    box(s, x, Inches(1.95), Inches(3.0), Inches(0.1), c)
    tb(s, x + Inches(0.15), Inches(2.2), Inches(2.7), Inches(0.4), t, 16, True, c)
    tb(s, x + Inches(0.15), Inches(2.7), Inches(2.7), Inches(1.3), b, 14, False, INK)

round_box(s, Inches(0.45), Inches(4.55), Inches(12.4), Inches(2.3), NAVY)
tb(s, Inches(0.7), Inches(4.75), Inches(12), Inches(0.4), "不变的五件套（扩端时禁止拆掉）", 15, True, TEAL)
tb(
    s,
    Inches(0.7),
    Inches(5.25),
    Inches(12),
    Inches(1.3),
    "1) YAML Case 与十个（或少量扩展的）原语    2) Execution Engine 与 Context 跨端传值\n3) Evidence / Reporter    4) Skill「禁止直连驱动」    5) SDK 作为唯一编程入口",
    15,
    False,
    WHITE,
)


# ========== 18 Extend how ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "扩到 App / iPad：建议工程步骤", "工作量在 Adapter 与定位策略，不在重写控制台")
footer(s, 25)
steps = [
    ("1. 定 target 名", "DSL 的 STEP_TARGETS 增加 app / ipad（或 ios、android）。一步只表达一个端，避免「mobile」大锅烩。"),
    ("2. 实现 TestAdapter", "新包 packages/adapter-app。把 click/input/assert 映射到底层：Appium、Maestro、Playwright Mobile、厂商私有工具皆可。"),
    ("3. 注册 Resolver", "引擎按 target 创建会话。一条 Case 仍可：app 下单 → extract → web 后台 assert。"),
    ("4. Locator 策略", "小程序已限制 css；App 可能要 accessibility id / 文案。在 Skill locator 规则里写清，禁止 AI 瞎编。"),
    ("5. 设备与环境", "testpilot.yaml 增加 app 段（包名、设备、Appium Server）。Workspace 绑「测试包 + 后台环境」。"),
    ("6. 证据与 Skill", "截图必做；系统弹窗/权限作为 wait/assert。Skill 增加 app 工作流，仍禁止直连 Appium。"),
]
for i, (t, b) in enumerate(steps):
    y = Inches(1.38) + Inches(i * 0.9)
    n = str(i + 1)
    round_box(s, Inches(0.45), y, Inches(12.4), Inches(0.82), WHITE, LINE)
    round_box(s, Inches(0.6), y + Inches(0.16), Inches(0.5), Inches(0.5), TEAL)
    tb(s, Inches(0.6), y + Inches(0.16), Inches(0.5), Inches(0.5), n, 14, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, Inches(1.3), y + Inches(0.08), Inches(3.2), Inches(0.65), t, 14, True, NAVY, MSO_ANCHOR.MIDDLE)
    tb(s, Inches(4.5), y + Inches(0.12), Inches(8.1), Inches(0.6), b, 13, False, MUTED)


# ========== FLOW 5 扩端决策流程 ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "流程图 ⑤　接入一个新端（App / iPad）的决策与工序", "左边是一次性的平台工作；右边是业务同事的日常，几乎不变")
footer(s)

round_box(s, Inches(0.4), Inches(1.35), Inches(7.6), Inches(5.5), WHITE, LINE)
tb(s, Inches(0.65), Inches(1.5), Inches(7.1), Inches(0.35), "平台组：一次性接入（估算以 Adapter 为主）", 14, True, NAVY)

fy = Inches(1.9)
node(s, Inches(0.7), fy, Inches(3.1), Inches(0.55), "确定端与 target 名：app / ipad", SOFT, NAVY, 12)
ar_d(s, Inches(2.16), fy + Inches(0.58), Inches(0.18), Inches(0.2))
diamond(s, Inches(0.95), Inches(2.72), Inches(2.6), Inches(0.95), "已有可用驱动?")
ar_r(s, Inches(3.65), Inches(3.11))
tb(s, Inches(4.1), Inches(2.72), Inches(3.6), Inches(0.95), "否 → 先做驱动选型（Appium / Maestro /\n厂商工具 / 设备云），产出一个能点通\n登录页的最小样例再开工。", 11, False, MUTED, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
ar_d(s, Inches(2.16), Inches(3.7), Inches(0.18), Inches(0.2))
ar_label(s, Inches(2.3), Inches(3.72), "是", TEAL)

flow_steps = [
    "实现 TestAdapter（packages/adapter-app）",
    "在 DSL 增加 target 枚举 + 注册 Resolver",
    "定 Locator 策略并写进 Skill 规则",
    "testpilot.yaml 增加设备/包配置与 Workspace",
    "补 Evidence（截图必做）+ 一条黄金 Case 验收",
]
for i, t in enumerate(flow_steps):
    y = Inches(3.98) + Inches(i * 0.55)
    node(s, Inches(0.7), y, Inches(7.0), Inches(0.44), t, WHITE, NAVY, 12, TEAL)
    if i < len(flow_steps) - 1:
        ar_d(s, Inches(4.11), y + Inches(0.45), Inches(0.16), Inches(0.09))

round_box(s, Inches(8.25), Inches(1.35), Inches(4.65), Inches(5.5), NAVY)
tb(s, Inches(8.5), Inches(1.5), Inches(4.2), Inches(0.35), "业务同事：几乎不用改习惯", 14, True, TEAL)
keep = [
    ("仍然写 YAML", "只是多了 target: app 这一种取值"),
    ("仍然 validate / run", "命令、报告、控制台都不变"),
    ("仍然一条 Case 跨端", "app 下单 → web 后台 assert"),
    ("存量用例不作废", "web / miniapp 的 Case 继续跑"),
    ("Skill 规则同步更新", "AI 照样只能写 DSL"),
]
for i, (t, b) in enumerate(keep):
    y = Inches(2.0) + Inches(i * 0.95)
    tb(s, Inches(8.5), y, Inches(4.2), Inches(0.35), "✓  " + t, 14, True, WHITE)
    tb(s, Inches(8.8), y + Inches(0.35), Inches(3.9), Inches(0.45), b, 12, False, RGBColor(0x9E, 0xD4, 0xD6))


# ========== 19 iPad specifics ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "App 与 iPad 不要当成「同一个手机端」", "屏幕、输入、多窗口会逼出 Adapter 的可选能力，而不是新 DSL")
footer(s, 26)
cards = [
    ("iPhone / Android 手机", "单窗口、手势（滑动/长按）可能要作为 Action 扩展或用现有 click+wait 组合。先覆盖核心路径，手势作为第二批原语。"),
    ("iPad", "分屏、Stage Manager、键盘外设。建议 target: ipad 独立，便于 Case 声明「仅大屏」。viewport 策略放配置，不写进步骤。"),
    ("混合技术栈", "原生 / Flutter / RN / H5 容器。Adapter 可对 H5 复用 Playwright webview，对原生走自动化树。对 DSL 仍是 click/assert。"),
    ("真机农场", "设备云是 Adapter 后面的基础设施（类似今天的微信开发者工具）。Control Plane 只记录哪次 Run 用了哪台设备快照。"),
]
for i, (t, b) in enumerate(cards):
    col, row = i % 2, i // 2
    x = Inches(0.45) + Inches(col * 6.45)
    y = Inches(1.5) + Inches(row * 2.6)
    round_box(s, x, y, Inches(6.25), Inches(2.4), WHITE, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.25), Inches(5.8), Inches(0.5), t, 16, True, TEAL)
    tb(s, x + Inches(0.25), y + Inches(0.85), Inches(5.8), Inches(1.3), b, 13, False, MUTED)


# ========== 20 Not doing ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "V0.1 明确不做：避免期望错位", "范围克制是优点。下列需求应另立项，而不是塞进 Core")
footer(s, 27)
nodo = [
    "RAG 知识库、从文档自动生成全部回归",
    "AI 自动改业务代码 / 自动修测试直到绿",
    "性能基线、业务数据基线、全链路监控替代品",
    "SQL/MySQL Adapter、把校验做成随意跑库",
    "业务 Skill（Login / Order / Payment）",
    "Dashboard 高级 IDE、可视化拖拽生成巨页",
    "Multi-Agent 分工编排、复杂任务调度器",
    "多用户权限系统、完整 SaaS 化账号体系",
]
for i, t in enumerate(nodo):
    col, row = i % 2, i // 2
    x = Inches(0.45) + Inches(col * 6.45)
    y = Inches(1.5) + Inches(row * 1.25)
    round_box(s, x, y, Inches(6.25), Inches(1.1), WHITE, LINE)
    tb(s, x + Inches(0.25), y + Inches(0.3), Inches(5.8), Inches(0.5), "✕   " + t, 15, False, INK)


# ========== FLOW 4 init 接入流程 ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "流程图 ④　npx csspilot init 在业务项目里做了什么", "全程幂等：已有文件一律保留，不动 .env、不动已有用例")
footer(s)

steps_a = [
    ("识别项目", "名称 / 类型 / 语言\n包管理器 / Git"),
    ("检测 Adapter", "Playwright、微信开发者工具\n缺失只告警，不阻塞"),
    ("生成配置", "testpilot.yaml\n.testpilot/project.json"),
    ("准备测试目录", "tests/e2e/{cases,fixtures,data}\n已存在则直接复用"),
]
steps_b = [
    ("安装 Skill", ".agents/skills/testpilot/\n+ references 项目上下文"),
    ("环境样例", ".env.example\n.gitignore 忽略 .env"),
    ("CI 提示", "检测已有 CI\n可用 ci init 生成模板"),
    ("自动 doctor", "体检并分组报告\nerror / warning"),
]
for i, (t, b) in enumerate(steps_a):
    x = Inches(0.45) + Inches(i * 3.2)
    round_box(s, x, Inches(1.45), Inches(2.85), Inches(1.5), WHITE, LINE)
    round_box(s, x + Inches(0.12), Inches(1.57), Inches(0.42), Inches(0.42), TEAL)
    tb(s, x + Inches(0.12), Inches(1.57), Inches(0.42), Inches(0.42), str(i + 1), 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.65), Inches(1.57), Inches(2.1), Inches(0.42), t, 14, True, NAVY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.15), Inches(2.08), Inches(2.6), Inches(0.8), b, 12, False, MUTED)
    if i < 3:
        ar_r(s, x + Inches(2.9), Inches(2.11))
ar_d(s, Inches(11.9), Inches(3.0), Inches(0.18), Inches(0.28))
tb(s, Inches(9.6), Inches(3.02), Inches(2.2), Inches(0.28), "接着装 AI 接入物", 11, True, TEAL, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

for i, (t, b) in enumerate(steps_b):
    x = Inches(0.45) + Inches(i * 3.2)
    round_box(s, x, Inches(3.45), Inches(2.85), Inches(1.5), WHITE, LINE)
    round_box(s, x + Inches(0.12), Inches(3.57), Inches(0.42), Inches(0.42), TEAL)
    tb(s, x + Inches(0.12), Inches(3.57), Inches(0.42), Inches(0.42), str(i + 5), 12, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.65), Inches(3.57), Inches(2.1), Inches(0.42), t, 14, True, NAVY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.15), Inches(4.08), Inches(2.6), Inches(0.8), b, 12, False, MUTED)
    if i < 3:
        ar_r(s, x + Inches(2.9), Inches(4.11))

round_box(s, Inches(0.45), Inches(5.2), Inches(6.1), Inches(1.65), WHITE, LINE)
tb(s, Inches(0.7), Inches(5.35), Inches(5.6), Inches(0.35), "init 之后业务项目多了什么", 14, True, TEAL)
tb(s, Inches(0.7), Inches(5.75), Inches(5.6), Inches(0.95), "testpilot.yaml  ·  .testpilot/  ·  tests/e2e/cases/\n.agents/skills/testpilot/  ·  .env.example", 13, False, MUTED)
round_box(s, Inches(6.8), Inches(5.2), Inches(6.1), Inches(1.65), NAVY)
tb(s, Inches(7.05), Inches(5.35), Inches(5.6), Inches(0.35), "绝不触碰", 14, True, GOLD)
tb(s, Inches(7.05), Inches(5.75), Inches(5.6), Inches(0.95), ".env 真实凭据  ·  已有用例  ·  已有 CI 配置\n升级用 csspilot update，只同步接入物", 13, False, WHITE)


# ========== 21 How to use ==========
s = prs.slides.add_slide(blank)
content_bg(s)
header_bar(s, "同事明天可以怎么用", "最小路径：接入 → 写/生成一条 Case → validate → run → 看报告")
footer(s, 28)
path = [
    ("接入", "在业务仓执行 npx csspilot init\n生成 Skill、testpilot.yaml、tests/e2e"),
    ("约定", "读 references/ 与 Skill 规则\n定位器必须有页面证据"),
    ("编写", "人写 YAML，或让 Agent 按工作流生成\n禁止让 AI 直接操作浏览器"),
    ("校验执行", "csspilot validate && run\nCI 里再加 sync-metadata"),
    ("看见", "本地 HTML 报告，或打开控制台\n看 Step、截图、失败原因"),
]
for i, (t, b) in enumerate(path):
    x = Inches(0.35) + Inches(i * 2.58)
    round_box(s, x, Inches(1.5), Inches(2.45), Inches(3.5), WHITE, LINE)
    round_box(s, x + Inches(0.85), Inches(1.7), Inches(0.7), Inches(0.7), TEAL)
    tb(s, x + Inches(0.85), Inches(1.7), Inches(0.7), Inches(0.7), str(i + 1), 18, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    tb(s, x + Inches(0.12), Inches(2.55), Inches(2.2), Inches(0.45), t, 16, True, NAVY, PP_ALIGN.CENTER)
    tb(s, x + Inches(0.12), Inches(3.1), Inches(2.2), Inches(1.6), b, 12, False, MUTED, PP_ALIGN.CENTER)

round_box(s, Inches(0.35), Inches(5.2), Inches(12.6), Inches(1.65), NAVY)
tb(s, Inches(0.6), Inches(5.4), Inches(12.1), Inches(0.35), "给负责人的投入建议", 14, True, TEAL)
tb(
    s,
    Inches(0.6),
    Inches(5.85),
    Inches(12.1),
    Inches(0.8),
    "先选 1 条跨端黄金路径（已有小程序下单→后台核验范式）在真实业务仓落地；同时约定「AI 只准写 DSL」。App/iPad 等有设备与驱动选型后再开 Adapter 专项，不必阻塞当前 Web/小程序收益。",
    14,
    False,
    WHITE,
)


# ========== 22 Close ==========
s = prs.slides.add_slide(blank)
box(s, 0, 0, W, H, NAVY)
box(s, 0, 0, Inches(0.18), H, TEAL)
tb(s, Inches(0.7), Inches(1.3), Inches(12), Inches(0.5), "带走的四句话", 14, False, TEAL)
lines = [
    "DSL 是 YAML 测试说明书；Adapter 是把它翻译成真实点击/请求的插件。",
    "仓库分层：入口（cli/web/agent）→ SDK → 引擎/DSL → Adapter。",
    "人和 AI 只写 DSL；真正碰端的只有 Adapter。",
    "App / iPad 是新 Adapter + 新 target，不是新测试平台。",
]
for i, line in enumerate(lines):
    y = Inches(1.95) + Inches(i * 0.7)
    tb(s, Inches(0.7), y, Inches(0.4), Inches(0.5), str(i + 1), 18, True, TEAL)
    tb(s, Inches(1.15), y, Inches(11.3), Inches(0.5), line, 18, False, WHITE)

tb(s, Inches(0.7), Inches(5.3), Inches(12), Inches(0.4), "讨论  ·  文档见 docs/architecture.md 与 README", 14, False, RGBColor(0x8A, 0xA0, 0xB8))
tb(s, Inches(0.7), Inches(5.85), Inches(12), Inches(0.4), "Q  &  A", 28, True, WHITE)

total_pages = len(prs.slides._sldIdLst)
for page, mark in PAGE_MARKS:
    mark.text_frame.paragraphs[0].runs[0].text = f"{page}  /  {total_pages}"

out = r"E:\css-test-pilot\docs\TestPilot-intro.pptx"
prs.save(out)
print("saved", out)
